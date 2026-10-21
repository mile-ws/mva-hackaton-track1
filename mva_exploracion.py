#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Exploración ordenada del dataset MVA Hackathon 2026.

Qué hace:
  1. Lista los archivos del dataset (con tamaño) y los clasifica.
  2. Baja el README y el documento del reto y los guarda como .txt para leerlos tranquila.
  3. Abre el VCF y guarda el ENCABEZADO (campos INFO, FORMAT, FILTER, cromosomas)
     en un Excel con una hoja por tema.
  4. Te dice si el VCF parece venir anotado y qué versión del genoma usa.

Todo se guarda en la carpeta:  Escritorio/mva_hackathon_salida

IMPORTANTE: por defecto NO se guardan filas de variantes (son datos del paciente).
Lo que sale en el Excel es solo metadata del archivo, no datos del niño.

Cómo se usa en Spyder:
  - Corré celda por celda con Shift+Enter (cada celda empieza con '# %%'),
    o todo junto con F5.
  - Necesitás estar logueada en Hugging Face (hf auth login) y tener acceso aprobado.
  - Si falta algo:  pip install openpyxl   (en la Terminal de la Mac)
"""

# %% 0. Configuración
from pathlib import Path

import pandas as pd
import pysam
from docx import Document
from huggingface_hub import HfApi, hf_hub_download

REPO_ID = "SageBio/mva-hackathon-2026-data"
REPO_TYPE = "dataset"

# El nombre del VCF no se escribe acá (lleva IDs de muestra del laboratorio y el código
# se publica): se toma del inventario en la celda 3.
ARCHIVO_README = "README.md"
ARCHIVO_RETO = "Challenge_Clinical_Phenotype_1.docx"

# Si lo ponés en True, se guarda una hoja con las primeras variantes del VCF.
# Son datos del paciente: quedan SOLO en tu compu, no los pases por chat
# ni los subas a GitHub, y recordá borrarlos al cerrar el hackatón.
GUARDAR_MUESTRA_VARIANTES = False
N_VARIANTES_MUESTRA = 200

CARPETA_SALIDA = Path.home() / "Desktop" / "mva_hackathon_salida"
CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)
print("Los resultados se guardan en:", CARPETA_SALIDA)


# %% 1. Inventario de archivos del dataset
def clasificar(nombre: str) -> str:
    n = nombre.lower()
    if n.endswith(".vcf.gz.tbi"):
        return "Índice VCF"
    if n.endswith(".vcf.gz"):
        return "VCF"
    if n.endswith((".fastq.gz", ".fq.gz")):
        return "FASTQ"
    if n.endswith((".bam", ".cram")):
        return "BAM/CRAM"
    if n.endswith(".docx"):
        return "Documento Word"
    if n.endswith(".md"):
        return "README"
    return "Otro"


api = HfApi()
filas = []
for item in api.list_repo_tree(repo_id=REPO_ID, repo_type=REPO_TYPE, recursive=True):
    tamano = getattr(item, "size", None)
    if tamano is None:  # es una carpeta
        continue
    filas.append(
        {
            "archivo": item.path,
            "tipo": clasificar(item.path),
            "tamaño_GB": round(tamano / 1e9, 3),
        }
    )

inventario = pd.DataFrame(filas).sort_values("tamaño_GB", ascending=False)
print(f"{len(inventario)} archivos. Tamaño total: {inventario['tamaño_GB'].sum():.1f} GB")
print(inventario.groupby("tipo")["tamaño_GB"].agg(["count", "sum"]).round(2))


# %% 2. README y documento del reto -> archivos de texto
ruta_readme = hf_hub_download(REPO_ID, ARCHIVO_README, repo_type=REPO_TYPE)
(CARPETA_SALIDA / "README.txt").write_text(
    Path(ruta_readme).read_text(encoding="utf-8"), encoding="utf-8"
)

ruta_reto = hf_hub_download(REPO_ID, ARCHIVO_RETO, repo_type=REPO_TYPE)
doc = Document(ruta_reto)
partes = [p.text for p in doc.paragraphs if p.text.strip()]
for tabla in doc.tables:  # por si hay datos clínicos en tablas
    for fila in tabla.rows:
        partes.append(" | ".join(c.text.strip() for c in fila.cells))
(CARPETA_SALIDA / "reto_fenotipo_clinico.txt").write_text(
    "\n".join(partes), encoding="utf-8"
)
print("Listo: README.txt y reto_fenotipo_clinico.txt guardados. Abrilos con calma.")


# %% 3. Bajar el VCF y su índice (si ya está en caché no lo vuelve a bajar)
vcfs = inventario.loc[inventario["tipo"] == "VCF", "archivo"].tolist()
if len(vcfs) != 1:
    raise RuntimeError(f"Esperaba un solo VCF en el dataset y encontré {len(vcfs)}.")
ruta_vcf = hf_hub_download(REPO_ID, vcfs[0], repo_type=REPO_TYPE)
ruta_tbi = hf_hub_download(REPO_ID, vcfs[0] + ".tbi", repo_type=REPO_TYPE)
vcf = pysam.VariantFile(ruta_vcf)
print("VCF abierto. Muestras:", list(vcf.header.samples))


# %% 4. Encabezado del VCF -> tablas
hdr = vcf.header

info = pd.DataFrame(
    [
        {"campo": k, "tipo": v.type, "número": str(v.number), "descripción": v.description}
        for k, v in hdr.info.items()
    ]
)
formato = pd.DataFrame(
    [
        {"campo": k, "tipo": v.type, "número": str(v.number), "descripción": v.description}
        for k, v in hdr.formats.items()
    ]
)
filtros = pd.DataFrame(
    [{"filtro": k, "descripción": v.description} for k, v in hdr.filters.items()]
)
contigs = pd.DataFrame(
    [{"cromosoma": c.name, "largo": c.length} for c in hdr.contigs.values()]
)
muestras = pd.DataFrame({"muestra": list(hdr.samples)})

print(f"Campos INFO: {len(info)} | FORMAT: {len(formato)} | contigs: {len(contigs)}")


# %% 5. Diagnóstico: ¿viene anotado? ¿qué versión del genoma?
pistas_anotacion = ["ANN", "CSQ", "BCSQ", "AF", "GNOMAD", "CLNSIG", "GENEINFO", "GENE"]
nombres_info = [c.upper() for c in info["campo"]]
encontradas = [p for p in pistas_anotacion if any(p in n for n in nombres_info)]

if encontradas:
    print("Posible anotación en INFO (campos que contienen):", encontradas)
else:
    print("No veo campos de anotación obvios: probablemente habrá que anotar aparte (VEP).")

# Versión del genoma según el largo del cromosoma 1
largo_chr1 = None
for nombre in ("chr1", "1"):
    if nombre in hdr.contigs:
        largo_chr1 = hdr.contigs[nombre].length
        break

if largo_chr1 == 248_956_422:
    version = "hg38 / GRCh38"
elif largo_chr1 == 249_250_621:
    version = "hg19 / GRCh37"
else:
    version = f"no se pudo determinar (largo chr1 = {largo_chr1})"
print("Versión del genoma:", version)
print("Nombres de cromosomas con prefijo 'chr':", "chr1" in hdr.contigs)


# %% 6. (Opcional) muestra de variantes, solo si lo activaste arriba
muestra_variantes = None
if GUARDAR_MUESTRA_VARIANTES:
    regs = []
    for i, r in enumerate(vcf.fetch()):
        regs.append(
            {
                "CHROM": r.chrom,
                "POS": r.pos,
                "ID": r.id,
                "REF": r.ref,
                "ALT": ",".join(r.alts) if r.alts else None,
                "QUAL": r.qual,
                "FILTER": ",".join(r.filter.keys()),
            }
        )
        if i + 1 >= N_VARIANTES_MUESTRA:
            break
    muestra_variantes = pd.DataFrame(regs)


# %% 7. Guardar todo en un Excel
ruta_excel = CARPETA_SALIDA / "mva_exploracion.xlsx"

resumen = pd.DataFrame(
    {
        "dato": [
            "Versión del genoma",
            "Nº de muestras en el VCF",
            "Campos INFO",
            "Campos FORMAT",
            "Posibles campos de anotación",
        ],
        "valor": [
            version,
            len(muestras),
            len(info),
            len(formato),
            ", ".join(encontradas) if encontradas else "ninguno obvio",
        ],
    }
)

hojas = {
    "resumen": resumen,
    "inventario": inventario,
    "INFO": info,
    "FORMAT": formato,
    "FILTER": filtros,
    "cromosomas": contigs,
    "muestras": muestras,
}
if muestra_variantes is not None:
    hojas["muestra_variantes"] = muestra_variantes

with pd.ExcelWriter(ruta_excel, engine="openpyxl") as writer:
    for nombre_hoja, tabla in hojas.items():
        tabla.to_excel(writer, sheet_name=nombre_hoja, index=False)
        ws = writer.sheets[nombre_hoja]
        for col in ws.columns:  # ancho de columnas legible
            ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(ancho + 2, 80)

print("Excel guardado en:", ruta_excel)
