#!/usr/bin/env python3
"""
mva_cruce_hpo_genes.py
Cruza una lista de términos HPO (síntomas) con el archivo público de HPO
(phenotype_to_genes.txt) y arma un ranking de genes candidatos.

El archivo de HPO es público. LO QUE GENERA ESTE SCRIPT (síntomas + genes del
caso) deriva de datos protegidos: no subirlo a GitHub, Drive ni mail. Ver
.gitignore sugerido al final del README del repo (carpeta de salida ignorada).

Cabecera verificada de phenotype_to_genes.txt (release descargada el 2026-10-20):
    hpo_id  hpo_name  ncbi_gene_id  gene_symbol  disease_id
Este archivo ya PROPAGA hacia arriba: un gen anotado a un término hijo
también figura en los términos padres. El script igual chequea la cabecera y
se detiene con un mensaje claro si cambió.

Uso mínimo:
    python mva_cruce_hpo_genes.py --hpo-ids mis_sintomas.csv \
        --annotations phenotype_to_genes.txt --out salida_hpo

Entrada de --hpo-ids: .csv/.txt con una columna "HPO ID" (o una lista de
códigos, uno por línea), o el .xlsx de la planilla (hoja "Síntomas HPO",
columna "HPO ID").
"""
import argparse
import csv
import math
import os
import re
import sys
from collections import defaultdict

EXPECTED_HEADER = ["hpo_id", "hpo_name", "ncbi_gene_id", "gene_symbol", "disease_id"]
HPO_RE = re.compile(r"^HP:\d{7}$")
MVA_GENES = ["BUB1B", "CEP57", "TRIP13", "CENATAC", "SLF2", "SMC5", "BUB1", "BUB3"]  # panel de 8 genes


# ----------------------------------------------------------------- lectura
def normalize_id(x):
    return str(x).strip().upper().replace(" ", "")


def read_hpo_ids(path):
    """Devuelve la lista de HPO IDs válidos (sin duplicados, en orden) y los inválidos."""
    raw = []
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Para leer .xlsx hace falta openpyxl (pip install openpyxl), o exportá la hoja a CSV.")
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb["Síntomas HPO"] if "Síntomas HPO" in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        col = next((i for i, h in enumerate(header) if h in ("hpo id", "hpo_id", "hpo")), None)
        if col is None:
            sys.exit("No encontré la columna 'HPO ID' en el Excel.")
        raw = [r[col] for r in rows[1:] if r[col]]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            sample = f.read(2048)
            f.seek(0)
            if "HP:" in sample.splitlines()[0].upper() and "," not in sample.splitlines()[0]:
                raw = [line.strip() for line in f if line.strip()]  # lista simple
            else:
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                except csv.Error:  # CSV de una sola columna: no hay separador que detectar
                    dialect = csv.excel
                reader = csv.reader(f, dialect)
                header = [h.strip().lower() for h in next(reader)]
                col = next((i for i, h in enumerate(header) if h in ("hpo id", "hpo_id", "hpo")), None)
                if col is None:
                    sys.exit("No encontré la columna 'HPO ID' en el CSV.")
                raw = [r[col] for r in reader if len(r) > col and r[col].strip()]
    ids, invalid, seen = [], [], set()
    for r in raw:
        n = normalize_id(r)
        if not HPO_RE.match(n):
            invalid.append(str(r))
        elif n not in seen:
            seen.add(n)
            ids.append(n)
    return ids, invalid


def read_gene_list(path):
    with open(path, encoding="utf-8-sig") as f:
        return {line.strip().upper() for line in f if line.strip() and not line.startswith("#")}


def read_annotations(path, wanted):
    """Una sola pasada: junta genes de los términos pedidos y cuenta genes totales."""
    term_genes = defaultdict(set)   # HP -> {símbolos}
    term_name = {}
    gene_id = {}                    # símbolo -> ncbi_gene_id
    all_genes = set()
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f, delimiter="\t")
        header = next(reader)
        if [h.strip() for h in header] != EXPECTED_HEADER:
            sys.exit(
                "La cabecera del archivo de HPO cambió.\n"
                f"  esperada: {EXPECTED_HEADER}\n  encontrada: {header}\n"
                "Ajustá EXPECTED_HEADER y los índices antes de seguir."
            )
        for r in reader:
            hp, name, ncbi, sym = r[0], r[1], r[2], r[3]
            all_genes.add(sym)
            if hp in wanted:
                term_genes[hp].add(sym)
                term_name[hp] = name
                gene_id[sym] = ncbi
    return term_genes, term_name, gene_id, len(all_genes)


# ----------------------------------------------------------------- cálculo
def weight_for(mode, n_term, n_total):
    if mode == "count":
        return 1.0
    if mode == "inverse":
        return 1.0 / n_term
    if mode == "idf":
        return math.log(n_total / n_term)
    raise ValueError(mode)


def build_ranking(ids, term_genes, gene_id, n_total, weight, max_genes, min_symptoms, only, mva):
    """Devuelve (filas del ranking, síntomas usados, síntomas excluidos).
    posicion = orden de la lista (desempate alfabético: NO significa nada dentro de un empate).
    puesto_con_empates = mejor puesto del grupo con el mismo puntaje (el que hay que informar)."""
    used, excluded = [], []
    for h in ids:
        if h not in term_genes:
            continue
        if max_genes is not None and len(term_genes[h]) > max_genes:
            excluded.append(h)
        else:
            used.append(h)
    per_gene = defaultdict(lambda: {"symptoms": [], "score": 0.0})
    for h in used:
        w = weight_for(weight, len(term_genes[h]), n_total)
        for g in term_genes[h]:
            per_gene[g]["symptoms"].append(h)
            per_gene[g]["score"] += w
    rows = []
    for g, d in per_gene.items():
        if len(d["symptoms"]) < min_symptoms:
            continue
        if only is not None and g.upper() not in only:
            continue
        rows.append({
            "gen": g, "ncbi_gene_id": gene_id.get(g, ""),
            "n_sintomas": len(d["symptoms"]),
            "pct_sintomas": round(100 * len(d["symptoms"]) / len(used), 1),
            "puntaje": round(d["score"], 4),
            "gen_MVA_conocido": "Sí" if g.upper() in mva else "No",
            "hpo_que_explica": ";".join(sorted(d["symptoms"])),
        })
    rows.sort(key=lambda r: (-r["puntaje"], -r["n_sintomas"], r["gen"]))
    # puestos con empates (mismo puntaje = mismo grupo)
    first_of = {}
    count_of = defaultdict(int)
    for i, r in enumerate(rows, start=1):
        r["posicion"] = i
        first_of.setdefault(r["puntaje"], i)
        count_of[r["puntaje"]] += 1
    for r in rows:
        r["puesto_con_empates"] = first_of[r["puntaje"]]
        r["genes_empatados"] = count_of[r["puntaje"]]
    return rows, used, excluded


def run_sensitivity(gene, ids, term_genes, gene_id, n_total, only, mva, min_symptoms, outdir):
    """Tabla configuración -> puesto del gen (con empates). Solo conteos y posiciones."""
    gene = gene.upper()
    out = []
    for weight in ("idf", "inverse", "count"):
        for cap in (None, 1000, 300):
            rows, used, excluded = build_ranking(ids, term_genes, gene_id, n_total, weight, cap,
                                                 min_symptoms, only, mva)
            rec = {"peso": weight, "max_genes_por_sintoma": cap if cap else "sin tope",
                   "sintomas_usados": len(used), "sintomas_excluidos": len(excluded),
                   "genes_en_ranking": len(rows)}
            hit = next((r for r in rows if r["gen"].upper() == gene), None)
            if not used:
                rec.update({"puesto_con_empates": "sin síntomas", "genes_empatados": "", "n_sintomas_del_gen": ""})
            elif hit is None:
                rec.update({"puesto_con_empates": "no aparece", "genes_empatados": "", "n_sintomas_del_gen": ""})
            else:
                rec.update({"puesto_con_empates": hit["puesto_con_empates"],
                            "genes_empatados": hit["genes_empatados"],
                            "n_sintomas_del_gen": hit["n_sintomas"]})
            out.append(rec)
    fields = ["peso", "max_genes_por_sintoma", "sintomas_usados", "sintomas_excluidos", "genes_en_ranking",
              "puesto_con_empates", "genes_empatados", "n_sintomas_del_gen"]
    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, f"sensibilidad_{gene}.csv")
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(out)
    print(f"\nSensibilidad de {gene} (puesto con empates = mejor puesto del grupo con el mismo puntaje):")
    print(f"{'peso':<8}{'tope genes':<12}{'síntomas':<10}{'genes rank.':<13}{'puesto':<12}{'empatados':<11}{'n_sint. gen'}")
    for r in out:
        print(f"{r['peso']:<8}{str(r['max_genes_por_sintoma']):<12}{r['sintomas_usados']:<10}{r['genes_en_ranking']:<13}"
              f"{str(r['puesto_con_empates']):<12}{str(r['genes_empatados']):<11}{r['n_sintomas_del_gen']}")
    print(f"Guardado en: {path}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--hpo-ids", required=True, help="CSV/TXT/XLSX con los HPO ID del caso")
    ap.add_argument("--annotations", required=True, help="phenotype_to_genes.txt de HPO")
    ap.add_argument("--out", default="salida_hpo", help="carpeta de salida (no subir a GitHub)")
    ap.add_argument("--weight", choices=["idf", "inverse", "count"], default="idf",
                    help="peso de cada síntoma: idf (default, log(N/n)), inverse (1/n) o count (todos 1)")
    ap.add_argument("--max-genes-per-symptom", type=int, default=None,
                    help="EXCLUIR del ranking los síntomas con más genes que este número (ej. 1500)")
    ap.add_argument("--min-symptoms", type=int, default=1,
                    help="solo genes que explican al menos N síntomas (default 1)")
    ap.add_argument("--only-genes", default=None,
                    help="archivo de texto con símbolos (uno por línea) para restringir el ranking, "
                         "por ejemplo los genes que tienen variantes candidatas tras el filtro del VCF")
    ap.add_argument("--extra-mva-genes", nargs="*", default=[], help="símbolos extra a marcar como 'gen MVA'")
    ap.add_argument("--top", type=int, default=0, help="guardar solo los N primeros del ranking (0 = todos)")
    ap.add_argument("--sensitivity-gene", default=None,
                    help="símbolo de un gen (ej. BUB1B): corre 9 configuraciones (3 pesos x tope sin límite/1000/300) "
                         "y guarda una tabla configuración -> puesto del gen")
    ap.add_argument("--xlsx", action="store_true", help="además de los CSV, escribir un Excel con las tablas")
    args = ap.parse_args()

    ids, invalid = read_hpo_ids(args.hpo_ids)
    if not ids:
        sys.exit("No quedó ningún HPO ID válido en la entrada.")
    if invalid:
        print(f"[aviso] {len(invalid)} entradas con formato inválido (se ignoran). Esperado: HP: + 7 dígitos.")

    term_genes, term_name, gene_id, n_total = read_annotations(args.annotations, set(ids))
    missing = [h for h in ids if h not in term_genes]
    if missing:
        print(f"[aviso] {len(missing)} HPO ID sin genes en el archivo (obsoleto, reemplazado o sin anotación): "
              f"{', '.join(missing)}")

    only = read_gene_list(args.only_genes) if args.only_genes else None
    mva = {g.upper() for g in MVA_GENES + args.extra_mva_genes}

    rows, used, excluded = build_ranking(ids, term_genes, gene_id, n_total, args.weight,
                                         args.max_genes_per_symptom, args.min_symptoms, only, mva)
    if not used:
        sys.exit("Ningún síntoma quedó después de los filtros. Revisá --max-genes-per-symptom.")
    if args.top:
        rows = rows[: args.top]

    # Tabla de síntomas (tamaño de cada lista)
    sym_rows = []
    for h in ids:
        n = len(term_genes.get(h, ()))
        estado = ("sin genes en el archivo" if h not in term_genes
                  else "EXCLUIDO (demasiados genes)" if h in excluded else "usado")
        sym_rows.append({"hpo_id": h, "hpo_name": term_name.get(h, ""), "n_genes": n,
                         "peso_por_gen": round(weight_for(args.weight, n, n_total), 4) if n else "",
                         "estado": estado})

    # Formato largo (para la hoja "Genes por síntoma" de la planilla)
    long_rows = [{"hpo_id": h, "gen": g, "ncbi_gene_id": gene_id.get(g, "")}
                 for h in used for g in sorted(term_genes[h])]

    # Salida
    os.makedirs(args.out, exist_ok=True)

    def write_csv(name, data, fields):
        with open(os.path.join(args.out, name), "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(data)

    write_csv("ranking_genes.csv", rows,
              ["posicion", "puesto_con_empates", "genes_empatados", "gen", "ncbi_gene_id", "n_sintomas",
               "pct_sintomas", "puntaje", "gen_MVA_conocido", "hpo_que_explica"])
    write_csv("sintomas_resumen.csv", sym_rows, ["hpo_id", "hpo_name", "n_genes", "peso_por_gen", "estado"])
    write_csv("genes_por_sintoma.csv", long_rows, ["hpo_id", "gen", "ncbi_gene_id"])

    with open(os.path.join(args.out, "parametros_usados.txt"), "w", encoding="utf-8") as f:
        f.write(f"archivo_anotaciones: {os.path.basename(args.annotations)}\n")
        f.write(f"genes_totales_en_archivo: {n_total}\n")
        f.write(f"sintomas_entrada: {len(ids)} | usados: {len(used)} | excluidos: {len(excluded)} | sin genes: {len(missing)}\n")
        f.write(f"peso: {args.weight}\nmax_genes_por_sintoma: {args.max_genes_per_symptom}\n")
        f.write(f"min_sintomas: {args.min_symptoms}\nrestringido_a_lista_de_genes: {bool(only)}\n")
        f.write(f"genes_en_ranking: {len(rows)}\n")
        f.write("(Guardar este archivo junto a la fecha de descarga de HPO para el informe de métodos.)\n")

    if args.xlsx:
        try:
            from openpyxl import Workbook
            wb = Workbook()
            for title, data, fields in (
                ("Ranking de genes", rows, ["posicion", "puesto_con_empates", "genes_empatados", "gen", "ncbi_gene_id",
                                            "n_sintomas", "pct_sintomas", "puntaje", "gen_MVA_conocido", "hpo_que_explica"]),
                ("Síntomas resumen", sym_rows, ["hpo_id", "hpo_name", "n_genes", "peso_por_gen", "estado"]),
                ("Genes por síntoma", long_rows, ["hpo_id", "gen", "ncbi_gene_id"]),
            ):
                ws = wb.active if title == "Ranking de genes" else wb.create_sheet()
                ws.title = title
                ws.append(fields)
                for r in data:
                    ws.append([r[k] for k in fields])
            wb.save(os.path.join(args.out, "resultado_hpo.xlsx"))
        except ImportError:
            print("[aviso] sin openpyxl: se omite el .xlsx (los CSV sí se escribieron).")

    # Resumen en consola (solo conteos y símbolos del ranking)
    print(f"\nSíntomas: {len(ids)} entrada | {len(used)} usados | {len(excluded)} excluidos | {len(missing)} sin genes")
    print(f"Genes en el ranking: {len(rows)}  (peso = {args.weight})")
    print("Top 10:")
    for r in rows[:10]:
        print(f"  {r['posicion']:>3}. {r['gen']:<10} síntomas={r['n_sintomas']:<3} puntaje={r['puntaje']:<8} "
              f"empate con {r['genes_empatados']-1} más | MVA={r['gen_MVA_conocido']}")
    mva_hits = [r for r in rows if r["gen_MVA_conocido"] == "Sí"]
    print("Genes MVA conocidos en el ranking:",
          ", ".join(f"{r['gen']} (puesto {r['puesto_con_empates']}, empatado con {r['genes_empatados']-1} más)" for r in mva_hits) or "ninguno")
    print(f"\nArchivos en: {args.out}/")
    if args.sensitivity_gene:
        run_sensitivity(args.sensitivity_gene, ids, term_genes, gene_id, n_total, only, mva,
                        args.min_symptoms, args.out)


if __name__ == "__main__":
    main()
